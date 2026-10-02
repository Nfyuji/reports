# -*- coding: utf-8 -*-
"""محرك قراءة وكتابة تقرير PowerPoint الفني."""

from __future__ import annotations

import copy
import io
import re
import shutil
import tempfile
from pathlib import Path
from typing import Any

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

NSMAP = {
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
    "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
    "a14": "http://schemas.microsoft.com/office/drawing/2010/main",
}

# مواضع صور القالب الأصلية (داخل الإطارات الزرقاء المزدوجة) — يمين→يسار ثم أعلى→أسفل
PHOTO_SLOTS = [
    (11833212, 1825602, 3323890, 3367144),
    (7615275, 1813175, 3305053, 3391998),
    (3276600, 1865376, 3376490, 3386911),
    (11800828, 6413630, 3446918, 3120568),
    (7615275, 6352006, 3305053, 3182192),
    (3366021, 6352006, 3252937, 3182192),
]

# مواضع الإطارات الخارجية (نفس ترتيب الخانات)
FRAME_SLOTS = [
    (11507945, 1483043, 3945588, 4052261),
    (7265277, 1483043, 3945588, 4052261),
    (3025014, 1483043, 3945588, 4052261),
    (11507945, 5947783, 3945588, 4052261),
    (7265277, 5947783, 3945588, 4052261),
    (3025014, 5947783, 3945588, 4052261),
]

PHOTOS_PER_SLIDE = 6
PHOTO_TITLE = " صور الملاحظات وصور عامة للمبنى "
ASSETS_DIR = Path(__file__).resolve().parent / "assets"
PHOTO_SEED_PATH = ASSETS_DIR / "photo_seed.pptx"


def find_report_file(folder: Path | None = None) -> Path:
    folder = folder or Path(__file__).resolve().parent
    files = sorted(folder.glob("*.pptx"), key=lambda p: p.stat().st_mtime, reverse=True)
    skip_prefixes = ("~$", "_")
    skip_names = {"_test_out.pptx"}
    files = [
        f
        for f in files
        if not f.name.startswith(skip_prefixes)
        and f.name not in skip_names
        and "_محدث" not in f.name
        and not f.name.endswith(".tmp.pptx")
    ]
    if not files:
        raise FileNotFoundError("لا يوجد ملف تقرير PowerPoint في المجلد.")
    # فضّل الملف الذي يحتوي «تقرير» في الاسم إن وُجد
    preferred = [f for f in files if "تقرير" in f.name]
    return preferred[0] if preferred else files[0]


def _set_run_text(run, text: str) -> None:
    run.text = text


def _parse_cover(prs: Presentation) -> dict[str, str]:
    slide = prs.slides[0]
    title = ""
    school = ""
    ministry_no = ""
    date = ""
    engineer = ""

    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if shape.name == "TextBox 2":
            title = shape.text_frame.paragraphs[0].runs[0].text if shape.text_frame.paragraphs[0].runs else shape.text
        elif shape.name == "TextBox 3":
            paras = shape.text_frame.paragraphs
            if len(paras) >= 1 and len(paras[0].runs) >= 3:
                school = paras[0].runs[2].text.strip()
            if len(paras) >= 2 and len(paras[1].runs) >= 3:
                ministry_no = paras[1].runs[2].text.strip()
            if len(paras) >= 3 and len(paras[2].runs) >= 2:
                date = paras[2].runs[1].text.strip()
            # سطر المهندس: إما تشغيل واحد «اسم المهندس :...» أو 3 تشغيلات
            for p in paras:
                joined = "".join(r.text for r in p.runs).strip()
                if "مهندس" not in joined and "المعد" not in joined:
                    continue
                if ":" in joined:
                    engineer = joined.split(":", 1)[1].strip()
                elif len(p.runs) >= 3:
                    engineer = p.runs[2].text.strip()
                break

    return {
        "title": title.strip() or "التقرير الفني لحالة المبنى",
        "school": school,
        "ministry_no": ministry_no,
        "date": date,
        "engineer": engineer,
    }


def _paragraph_plain_text(p_el) -> str:
    """يجمع نص الفقرة بما فيه أجزاء المعادلات إن وُجدت."""
    parts: list[str] = []
    for t_el in p_el.xpath(".//a:t | .//m:t", namespaces={
        **NSMAP,
        "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
    }):
        if t_el.text:
            parts.append(t_el.text)
    return "".join(parts).replace("\u00a0", " ").strip()


def _extract_observations_from_xml(slide_xml: bytes) -> list[str]:
    root = etree.fromstring(slide_xml)
    notes: list[str] = []

    paragraphs = root.xpath(
        './/mc:Choice//p:sp[p:nvSpPr/p:cNvPr[@name="مربع نص 4"]]/p:txBody/a:p',
        namespaces=NSMAP,
    )
    if not paragraphs:
        paragraphs = root.xpath(
            './/p:sp[p:nvSpPr/p:cNvPr[@name="مربع نص 4"]]/p:txBody/a:p',
            namespaces=NSMAP,
        )

    for p in paragraphs:
        text = _paragraph_plain_text(p)
        if text:
            notes.append(text)

    if notes:
        return notes

    # احتياطي
    for t_el in root.xpath(".//a:t", namespaces=NSMAP):
        text = (t_el.text or "").strip()
        if text and text not in {"إجمالي الملاحظات", "\u00a0"}:
            notes.append(text)
    return notes


def _obs_paragraph_xml(text: str) -> etree._Element:
    """فقرة ملاحظة بنفس تنسيق القالب (رصاصة Wingdings + Tajawal)."""
    p = etree.Element(qn("a:p"))
    pPr = etree.SubElement(p, qn("a:pPr"))
    pPr.set("marL", "914400")
    pPr.set("lvl", "1")
    pPr.set("indent", "-457200")
    pPr.set("algn", "r")
    pPr.set("rtl", "1")
    buFont = etree.SubElement(pPr, qn("a:buFont"))
    buFont.set("typeface", "Wingdings")
    buFont.set("panose", "05000000000000000000")
    buFont.set("pitchFamily", "2")
    buFont.set("charset", "2")
    buChar = etree.SubElement(pPr, qn("a:buChar"))
    buChar.set("char", "§")

    r = etree.SubElement(p, qn("a:r"))
    rPr = etree.SubElement(r, qn("a:rPr"))
    rPr.set("lang", "ar-SA")
    rPr.set("sz", "4165")
    rPr.set("dirty", "0")
    solid = etree.SubElement(rPr, qn("a:solidFill"))
    etree.SubElement(solid, qn("a:srgbClr")).set("val", "083144")
    etree.SubElement(rPr, qn("a:latin")).set("typeface", "Tajawal")
    etree.SubElement(rPr, qn("a:ea")).set("typeface", "Tajawal")
    etree.SubElement(rPr, qn("a:cs")).set("typeface", "Tajawal")
    etree.SubElement(rPr, qn("a:rtl"))
    t = etree.SubElement(r, qn("a:t"))
    t.text = text
    return p


def _find_notes_txbody(root):
    """يجد txBody لمربع الملاحظات سواء كان العنصر oxml أو lxml."""
    # جرّب xpath بأسلوب python-pptx (بدون namespaces=)
    try:
        found = root.xpath('.//p:sp[p:nvSpPr/p:cNvPr[@name="مربع نص 4"]]/p:txBody')
        if found:
            return found
    except Exception:
        pass

    try:
        found = root.xpath(
            './/mc:Choice//p:sp[p:nvSpPr/p:cNvPr[@name="مربع نص 4"]]/p:txBody',
            namespaces=NSMAP,
        )
        if found:
            return found
        found = root.xpath(
            './/p:sp[p:nvSpPr/p:cNvPr[@name="مربع نص 4"]]/p:txBody',
            namespaces=NSMAP,
        )
        if found:
            return found
    except TypeError:
        pass

    # بحث يدوي
    results = []
    for sp in root.iter(qn("p:sp")):
        nv = sp.find(qn("p:nvSpPr"))
        if nv is None:
            continue
        cNvPr = nv.find(qn("p:cNvPr"))
        if cNvPr is not None and cNvPr.get("name") == "مربع نص 4":
            tx = sp.find(qn("p:txBody"))
            if tx is not None:
                results.append(tx)
    return results


def _write_observations_xml(slide_part, observations: list[str]) -> None:
    """يستبدل فقرات الملاحظات داخل مربع النص (AlternateContent)."""
    # أعِد التحليل بـ lxml لضمان تعديل موثوق ثم أعد الحقن في الجزء
    xml_bytes = slide_part.blob
    root = etree.fromstring(xml_bytes)

    tx_bodies = _find_notes_txbody(root)
    if not tx_bodies:
        raise RuntimeError("تعذر العثور على مربع الملاحظات في الشريحة الثانية.")

    clean = [o.strip() for o in observations if o and o.strip()]
    if not clean:
        clean = ["لا توجد ملاحظات"]

    for txBody in tx_bodies:
        for child in list(txBody):
            if etree.QName(child).localname == "p":
                txBody.remove(child)
        for note in clean:
            txBody.append(_obs_paragraph_xml(note))

    new_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    # حقن العنصر المحدّث في الجزء
    from pptx.oxml import parse_xml

    slide_part._element = parse_xml(new_xml)


def _is_photo_slide(slide) -> bool:
    pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    return len(pics) >= 1


def _picture_shapes(slide):
    """صور المحتوى فقط (تجاهل الشعارات الصغيرة إن أمكن)."""
    pics = []
    for s in slide.shapes:
        if s.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        # صور التقرير تقريباً مربعة وكبيرة
        if s.width and s.height and int(s.width) > 2000000 and int(s.height) > 2000000:
            pics.append(s)
            continue
        if "صورة" in (s.name or "") or "Picture" in (s.name or ""):
            if s.width and int(s.width) > 1500000:
                pics.append(s)
    return pics


def _delete_shape(shape) -> None:
    sp = shape._element
    sp.getparent().remove(sp)


def _duplicate_slide(prs: Presentation, index: int):
    """ينسخ شريحة مع علاقاتها بنفس معرفات rId (مهم للصور والزخارف/الإطارات)."""
    source = prs.slides[index]
    dest = prs.slides.add_slide(source.slide_layout)

    for el in list(dest.shapes._spTree):
        tag = etree.QName(el).localname
        if tag not in {"nvGrpSpPr", "grpSpPr"}:
            dest.shapes._spTree.remove(el)

    for el in source.shapes._spTree:
        tag = etree.QName(el).localname
        if tag in {"nvGrpSpPr", "grpSpPr"}:
            continue
        dest.shapes._spTree.append(copy.deepcopy(el))

    for rel in source.part.rels.values():
        if "slideLayout" in rel.reltype:
            continue
        try:
            dest.part.relate_to(rel._target, rel.reltype, rId=rel.rId)
        except Exception:
            try:
                dest.part.relate_to(rel._target, rel.reltype)
            except Exception:
                pass

    return dest


def _delete_slide(prs: Presentation, index: int) -> None:
    sldIdLst = prs.slides._sldIdLst
    slides = list(sldIdLst)
    sldId = slides[index]
    rId = sldId.get(qn("r:id"))
    sldIdLst.remove(sldId)
    prs.part.drop_rel(rId)


def _read_picture_bytes(picture) -> bytes:
    return picture.image.blob


def load_report(pptx_path: Path | None = None) -> dict[str, Any]:
    path = Path(pptx_path) if pptx_path else find_report_file()
    prs = Presentation(str(path))

    cover = _parse_cover(prs)
    observations = _extract_observations_from_xml(prs.slides[1].part.blob)

    photos: list[dict[str, Any]] = []
    for si, slide in enumerate(prs.slides):
        if si < 2:
            continue
        for pi, pic in enumerate(_sorted_content_pictures(slide)):
            blob = _read_picture_bytes(pic)
            photos.append(
                {
                    "slide": si + 1,
                    "index": pi,
                    "name": pic.name,
                    "blob": blob,
                    "content_type": pic.image.content_type,
                }
            )

    return {
        "path": str(path),
        "filename": path.name,
        "cover": cover,
        "observations": observations,
        "photos": photos,
        "photo_count": len(photos),
        "slide_count": len(prs.slides),
    }


def _ensure_engineer_paragraph(shape, engineer: str) -> None:
    """
    يضيف سطر المهندس بنسخ فقرة المدرسة حرفياً، لكن كنص واحد في تشغيل واحد
    حتى لا يعكس PowerPoint الاسم بسبب خلط اتجاه التشغيلات.
    الشكل النهائي: «اسم المهندس :ضاوي» بنفس اتجاه باقي الحقول.
    """
    tf = shape.text_frame
    value = (engineer or "").strip()
    line = f"اسم المهندس :{value}" if value else ""

    # احذف أي سطر مهندس قديم (حتى لو باتجاه خاطئ)
    for p in list(tf.paragraphs):
        joined = "".join(r.text for r in p.runs)
        if "مهندس" in joined or "المعد" in joined:
            el = p._p
            parent = el.getparent()
            if parent is not None:
                parent.remove(el)

    if not value or not tf.paragraphs:
        return

    template_p = tf.paragraphs[0]._p
    new_p = copy.deepcopy(template_p)

    # اجعلها تشغيلاً واحداً فقط (انسخ تنسيق أول run)
    runs = new_p.findall(qn("a:r"))
    if not runs:
        return

    first = runs[0]
    t = first.find(qn("a:t"))
    if t is None:
        t = etree.SubElement(first, qn("a:t"))
    t.text = line

    # تأكد من rtl على التشغيل
    rPr = first.find(qn("a:rPr"))
    if rPr is None:
        rPr = etree.SubElement(first, qn("a:rPr"))
    if rPr.find(qn("a:rtl")) is None:
        etree.SubElement(rPr, qn("a:rtl"))
    rPr.set("lang", "ar-SA")

    for extra in runs[1:]:
        new_p.remove(extra)

    # تأكد من rtl على الفقرة
    pPr = new_p.find(qn("a:pPr"))
    if pPr is None:
        pPr = etree.Element(qn("a:pPr"))
        new_p.insert(0, pPr)
    pPr.set("algn", "r")
    pPr.set("rtl", "1")

    tf._txBody.append(new_p)


def _update_cover(prs: Presentation, cover: dict[str, str]) -> None:
    slide = prs.slides[0]
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        if shape.name == "TextBox 2":
            if shape.text_frame.paragraphs[0].runs:
                _set_run_text(shape.text_frame.paragraphs[0].runs[0], cover.get("title", "").strip())
        elif shape.name == "TextBox 3":
            paras = shape.text_frame.paragraphs
            if len(paras) >= 1 and len(paras[0].runs) >= 3:
                _set_run_text(paras[0].runs[2], cover.get("school", "").strip())
            if len(paras) >= 2 and len(paras[1].runs) >= 3:
                _set_run_text(paras[1].runs[2], cover.get("ministry_no", "").strip())
            if len(paras) >= 3:
                if len(paras[2].runs) >= 2:
                    _set_run_text(paras[2].runs[1], cover.get("date", "").strip())
                elif len(paras[2].runs) == 1:
                    _set_run_text(paras[2].runs[0], f"التاريخ:{cover.get('date', '').strip()}")
            _ensure_engineer_paragraph(shape, cover.get("engineer", ""))


def _update_photo_title(slide, title: str = PHOTO_TITLE) -> None:
    for shape in slide.shapes:
        if shape.has_text_frame and "صور" in (shape.text or ""):
            if shape.text_frame.paragraphs and shape.text_frame.paragraphs[0].runs:
                shape.text_frame.paragraphs[0].runs[0].text = title
            break


def _sorted_content_pictures(slide):
    """يرتب صور المحتوى يمين→يسار ثم أعلى→أسفل (نفس ترتيب القالب)."""
    pics = _picture_shapes(slide)
    return sorted(
        pics,
        key=lambda p: (0 if int(p.top) < 4_500_000 else 1, -int(p.left)),
    )


def _replace_picture_image(picture, image_bytes: bytes) -> None:
    """يستبدل صورة الخانة بصورة جديدة مع الإبقاء على الموضع/الحجم/الترتيب (والإطار حوله)."""
    part = picture.part
    _image_part, rId = part.get_or_add_image_part(io.BytesIO(image_bytes))
    blip = picture._element.blipFill.blip
    try:
        blip.rEmbed = rId
    except Exception:
        blip.set(qn("r:embed"), rId)


def _ensure_picture_slots(slide, count: int = PHOTOS_PER_SLIDE) -> list:
    """
    يضمن وجود خانات صور بعدد count في مواضع القالب.
    الإطارات الزرقاء موجودة في Group 2؛ الصور تبقى داخلها.
    """
    slots = _sorted_content_pictures(slide)
    if len(slots) >= count:
        return slots[:count]

    existing_keys = {(int(p.left), int(p.top)) for p in slots}
    from PIL import Image as PILImage

    for left, top, width, height in PHOTO_SLOTS:
        if len(_sorted_content_pictures(slide)) >= count:
            break
        if any(abs(left - x) < 80_000 and abs(top - y) < 80_000 for x, y in existing_keys):
            continue
        buf = io.BytesIO()
        PILImage.new("RGB", (8, 8), (255, 255, 255)).save(buf, format="JPEG")
        buf.seek(0)
        pic = slide.shapes.add_picture(buf, Emu(left), Emu(top), width=Emu(width), height=Emu(height))
        existing_keys.add((left, top))
        slots.append(pic)

    return _sorted_content_pictures(slide)[:count]


def _frame_groups_sorted(slide):
    """
    إطارات الصور الزرقاء داخل Group 2 مرتبة بنفس ترتيب الخانات
    (يمين→يسار، أعلى→أسفل).
    """
    root = slide.part._element
    try:
        groups = root.xpath('.//p:grpSp[p:nvGrpSpPr/p:cNvPr[@name="Group 2"]]')
    except Exception:
        groups = []
    if not groups:
        return []

    g2 = groups[0]
    try:
        xfrm = g2.xpath("./p:grpSpPr/a:xfrm")[0]
        off = xfrm.xpath("./a:off")[0]
        ext = xfrm.xpath("./a:ext")[0]
        ch_ext = xfrm.xpath("./a:chExt")[0]
        gx, gy = int(off.get("x")), int(off.get("y"))
        gcx, gcy = int(ext.get("cx")), int(ext.get("cy"))
        chcx, chcy = int(ch_ext.get("cx")), int(ch_ext.get("cy"))
    except Exception:
        return []

    frames = []
    for child in list(g2):
        if etree.QName(child).localname != "grpSp":
            continue
        # الإطار الأزرق المزدوج فقط (تجاهل الزخارف الصغيرة)
        has_blue = False
        for clr in child.iter(qn("a:srgbClr")):
            if clr.get("val", "").upper() == "1A345B":
                has_blue = True
                break
        if not has_blue:
            continue
        try:
            cxfrm = child.xpath("./p:grpSpPr/a:xfrm")[0]
            coff = cxfrm.xpath("./a:off")[0]
            x = int(coff.get("x"))
            y = int(coff.get("y"))
            abs_x = gx + int(x * gcx / chcx)
            abs_y = gy + int(y * gcy / chcy)
        except Exception:
            continue
        frames.append((abs_x, abs_y, child))

    frames.sort(key=lambda t: (0 if t[1] < 4_500_000 else 1, -t[0]))
    return [el for _x, _y, el in frames]


def _add_double_frame(slide, index: int, left: int, top: int, width: int, height: int) -> None:
    """إطار أزرق مزدوج بزوايا مستديرة حول خانة الصورة."""
    for j, (inset, width_pt) in enumerate(((0, 3.0), (90_000, 1.5))):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Emu(left + inset),
            Emu(top + inset),
            Emu(width - 2 * inset),
            Emu(height - 2 * inset),
        )
        shape.name = f"إطار_{index}_{j}"
        shape.fill.background()
        shape.line.color.rgb = RGBColor(0x1A, 0x34, 0x5B)
        shape.line.width = Pt(width_pt)
        try:
            shape.adjustments[0] = 0.08
        except Exception:
            pass


def _build_photo_slide(prs: Presentation):
    """
    ينشئ شريحة صور جديدة (عند فقدان شرائح الصور من الملف):
    ينسخ زخارف الغلاف + يضيف العنوان + 6 إطارات + 6 خانات صور.
    """
    from PIL import Image as PILImage

    slide = _duplicate_slide(prs, 0)
    # احذف نصوص الغلاف فقط
    for sh in list(slide.shapes):
        if sh.has_text_frame and sh.name in {"TextBox 2", "TextBox 3"}:
            _delete_shape(sh)

    title_box = slide.shapes.add_textbox(Emu(6639744), Emu(304780), Emu(9497375), Emu(616836))
    title_box.name = "عنوان الصور"
    p = title_box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run = p.add_run()
    run.text = PHOTO_TITLE.strip()
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0x08, 0x31, 0x44)
    run.font.name = "Tajawal"

    for i, ((fl, ft, fw, fh), (pl, pt, pw, ph)) in enumerate(zip(FRAME_SLOTS, PHOTO_SLOTS)):
        _add_double_frame(slide, i, fl, ft, fw, fh)
        buf = io.BytesIO()
        PILImage.new("RGB", (48, 48), (245, 245, 245)).save(buf, format="JPEG")
        buf.seek(0)
        pic = slide.shapes.add_picture(buf, Emu(pl), Emu(pt), width=Emu(pw), height=Emu(ph))
        pic.name = f"صورة {i + 1}"

    return slide


def _save_photo_seed(prs: Presentation) -> None:
    """يحفظ شريحة صور كبذرة احتياطية حتى لا تُفقد لاحقاً."""
    if len(prs.slides) < 3:
        return
    try:
        ASSETS_DIR.mkdir(exist_ok=True)
        seed = Presentation()
        seed.slide_width = prs.slide_width
        seed.slide_height = prs.slide_height
        # أنشئ شريحة فارغة ثم انسخ محتوى شريحة الصور الأولى
        blank = seed.slide_layouts[6] if len(seed.slide_layouts) > 6 else seed.slide_layouts[0]
        dest = seed.slides.add_slide(blank)
        source = prs.slides[2]
        for el in list(dest.shapes._spTree):
            tag = etree.QName(el).localname
            if tag not in {"nvGrpSpPr", "grpSpPr"}:
                dest.shapes._spTree.remove(el)
        for el in source.shapes._spTree:
            tag = etree.QName(el).localname
            if tag in {"nvGrpSpPr", "grpSpPr"}:
                continue
            dest.shapes._spTree.append(copy.deepcopy(el))
        for rel in source.part.rels.values():
            if "slideLayout" in rel.reltype:
                continue
            try:
                dest.part.relate_to(rel._target, rel.reltype, rId=rel.rId)
            except Exception:
                try:
                    dest.part.relate_to(rel._target, rel.reltype)
                except Exception:
                    pass
        seed.save(str(PHOTO_SEED_PATH))
    except Exception:
        pass


def _ensure_photo_template(prs: Presentation) -> None:
    """يضمن وجود شريحة صور واحدة على الأقل (يبنيها إن كانت مفقودة)."""
    if len(prs.slides) >= 3:
        return
    _build_photo_slide(prs)
    _save_photo_seed(prs)


def _named_frame_indices(slide) -> dict[int, list]:
    """إطارات مُعاد بناؤها باسم إطار_i_j."""
    by_idx: dict[int, list] = {}
    for sh in slide.shapes:
        name = sh.name or ""
        if not name.startswith("إطار_"):
            continue
        parts = name.split("_")
        if len(parts) < 2:
            continue
        try:
            idx = int(parts[1])
        except ValueError:
            continue
        by_idx.setdefault(idx, []).append(sh._element)
    return by_idx


def _remove_unused_frames(slide, keep_count: int) -> None:
    """يحذف إطارات الخانات غير المستخدمة فقط."""
    # إطارات القالب الأصلية (Group 2)
    frames = _frame_groups_sorted(slide)
    for i, frame_el in enumerate(frames):
        if i >= keep_count:
            parent = frame_el.getparent()
            if parent is not None:
                parent.remove(frame_el)

    # إطارات مُعاد بناؤها
    named = _named_frame_indices(slide)
    for idx, elements in named.items():
        if idx >= keep_count:
            for el in elements:
                parent = el.getparent()
                if parent is not None:
                    parent.remove(el)


def _fill_slide_photos(slide, image_items: list[tuple[bytes, str]], trim_frames: bool = True) -> None:
    """
    يملأ شريحة الصور بعدد الصور فقط:
    - يستبدل الخانات المستخدمة ويُبقي إطارها
    - يحذف الخانات والإطارات الزائدة عند trim_frames=True
    """
    _update_photo_title(slide)
    n = len(image_items)
    slots = _ensure_picture_slots(slide, PHOTOS_PER_SLIDE)

    for i, pic in enumerate(list(slots)):
        if i < n:
            _replace_picture_image(pic, image_items[i][0])
        else:
            _delete_shape(pic)

    if trim_frames:
        _remove_unused_frames(slide, n)


def _guess_ext(filename: str | None, content_type: str | None) -> str:
    if filename:
        suf = Path(filename).suffix.lower()
        if suf in {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif"}:
            return ".jpg" if suf == ".jpeg" else suf
    if content_type:
        if "png" in content_type:
            return ".png"
        if "webp" in content_type:
            return ".webp"
    return ".jpg"


def save_report(
    pptx_path: Path,
    cover: dict[str, str],
    observations: list[str],
    photo_blobs: list[tuple[bytes, str]],
    output_path: Path | None = None,
) -> Path:
    """
    يعدّل ملف التقرير ويحفظه.
    عدد شرائح الصور = ceil(عدد الصور / 6).
    لا يحذف شريحة القالب الاحتياطية نهائياً حتى لو عدد الصور = 0.
    """
    pptx_path = Path(pptx_path)
    output_path = Path(output_path) if output_path else pptx_path

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp) / "work.pptx"
        shutil.copy2(pptx_path, tmp_path)
        prs = Presentation(str(tmp_path))

        _update_cover(prs, cover)
        if len(prs.slides) < 2:
            raise RuntimeError("ملف التقرير ناقص (يحتاج شريحة غلاف + ملاحظات على الأقل).")
        _write_observations_xml(prs.slides[1].part, observations)

        # إن كانت شرائح الصور مفقودة (بسبب حفظ قديم)، أعد بناءها
        _ensure_photo_template(prs)

        photo_indices = list(range(2, len(prs.slides)))
        template_idx = photo_indices[0]
        needed = (len(photo_blobs) + PHOTOS_PER_SLIDE - 1) // PHOTOS_PER_SLIDE if photo_blobs else 0

        if needed == 0:
            # أبقِ شريحة صور واحدة كقالب (بدون صور) حتى لا يضيع التصميم
            while len(photo_indices) > 1:
                _delete_slide(prs, photo_indices[-1])
                photo_indices = list(range(2, len(prs.slides)))
            seed_slide = prs.slides[photo_indices[0]]
            # امسح الصور فقط، أبقِ الإطارات كاملة للمرة القادمة
            for pic in list(_sorted_content_pictures(seed_slide)):
                _delete_shape(pic)
            _update_photo_title(seed_slide)
            _save_photo_seed(prs)
        else:
            # انسخ الشرائح الإضافية أولاً من القالب قبل التعبئة
            while len(photo_indices) < needed:
                _duplicate_slide(prs, template_idx)
                photo_indices = list(range(2, len(prs.slides)))

            while len(photo_indices) > needed:
                _delete_slide(prs, photo_indices[-1])
                photo_indices = list(range(2, len(prs.slides)))

            for page, slide_idx in enumerate(photo_indices):
                slide = prs.slides[slide_idx]
                start = page * PHOTOS_PER_SLIDE
                chunk = photo_blobs[start : start + PHOTOS_PER_SLIDE]
                # على الشريحة الأولى قبل النسخ كنا نحتاج إطارات كاملة —
                # بعد النسخ: احذف الزائد في كل شريحة حسب عدد صورها
                _fill_slide_photos(slide, chunk, trim_frames=True)

            _save_photo_seed(prs)

        prs.save(str(tmp_path))

        final_tmp = output_path.with_suffix(".tmp.pptx")
        shutil.copy2(tmp_path, final_tmp)
        try:
            final_tmp.replace(output_path)
        except PermissionError:
            alt = output_path.with_name(output_path.stem + "_محدث" + output_path.suffix)
            final_tmp.replace(alt)
            return alt

    return output_path

def photos_to_data_uris(photos: list[dict[str, Any]], max_preview: int = 120) -> list[dict[str, str]]:
    """للواجهة: يحول الصور إلى data URI مصغّرة عبر Pillow إن أمكن."""
    import base64

    try:
        from PIL import Image
    except ImportError:
        Image = None

    out = []
    for i, ph in enumerate(photos[:max_preview]):
        blob = ph["blob"]
        ctype = ph.get("content_type") or "image/jpeg"
        if Image is not None:
            try:
                im = Image.open(io.BytesIO(blob))
                im.thumbnail((360, 360))
                buf = io.BytesIO()
                if im.mode in ("RGBA", "P"):
                    im = im.convert("RGB")
                im.save(buf, format="JPEG", quality=72)
                blob = buf.getvalue()
                ctype = "image/jpeg"
            except Exception:
                pass
        b64 = base64.b64encode(blob).decode("ascii")
        out.append(
            {
                "id": str(ph.get("id", i)),
                "src": f"data:{ctype};base64,{b64}",
                "name": ph.get("name") or f"صورة {i + 1}",
            }
        )
    return out
